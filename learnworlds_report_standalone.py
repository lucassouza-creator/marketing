#!/usr/bin/env python3
"""
LearnWorlds → Excel Report Generator
=====================================
Busca dados do curso na LearnWorlds e gera um arquivo .xlsx
com 7 abas de análise pronto para subir no Google Sheets.

Dependências (instale uma vez):
    pip3 install requests python-dotenv pandas openpyxl

Uso:
    python3 learnworlds_report_standalone.py formacao-analise-dados-corteva
"""

import sys, os, time
from datetime import datetime, timezone
from pathlib import Path

from dotenv import load_dotenv
load_dotenv()

import requests
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import ColorScaleRule

# ── Credenciais ───────────────────────────────────────────────────────────────

BASE_URL      = os.getenv("LEARNWORLDS_URL", "https://academy.indicium.ai/admin/api")
CLIENT_ID     = os.getenv("LEARNWORLDS_CLIENT_ID")
CLIENT_SECRET = os.getenv("LEARNWORLDS_CLIENT_SECRET")

# ── Helpers de estilo ─────────────────────────────────────────────────────────

def fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def font(color="000000", bold=False, size=10, italic=False):
    return Font(color=color, bold=bold, size=size, italic=italic, name="Calibri")

def align(h="left", v="center", wrap=False):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

NAVY      = "1E3A5F"
WHITE     = "FFFFFF"
LT_GRAY   = "F2F4F7"
GREEN_BG  = "D4EDDA"; GREEN_FG  = "155724"
AMBER_BG  = "FFF3CD"; AMBER_FG  = "856404"
RED_BG    = "F8D7DA"; RED_FG    = "721C24"
GRAY_BG   = "E9ECEF"; GRAY_FG   = "495057"

STATUS_PT = {
    "completed":   "Concluído",
    "in_progress": "Em Andamento",
    "not_started": "Não Iniciado",
}
STATUS_COLORS = {
    "Concluído":    (GREEN_BG, GREEN_FG),
    "Em Andamento": (AMBER_BG, AMBER_FG),
    "Não Iniciado": (RED_BG,   RED_FG),
    "Desconhecido": (GRAY_BG,  GRAY_FG),
}

def _status_fill(status_pt):
    return STATUS_COLORS.get(status_pt, (GRAY_BG, GRAY_FG))

def _access_label(days):
    if days is None or (isinstance(days, float) and days != days):
        return "Nunca acessou", RED_BG, RED_FG
    if days <= 7:   return "Ativo (≤7 dias)",     GREEN_BG, GREEN_FG
    if days <= 30:  return "Recente (8-30 dias)",  AMBER_BG, AMBER_FG
    return "Inativo (>30 dias)", RED_BG, RED_FG

def _header_row(ws, row, values, bg=NAVY, fg=WHITE, height=22):
    ws.row_dimensions[row].height = height
    for col, val in enumerate(values, 1):
        c = ws.cell(row=row, column=col, value=val)
        c.fill = fill(bg); c.font = font(fg, bold=True)
        c.alignment = align(h="center")

def _col_widths(ws, widths: dict):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

def _color_scale(ws, min_row, max_row, col_letter):
    ws.conditional_formatting.add(
        f"{col_letter}{min_row}:{col_letter}{max_row}",
        ColorScaleRule(
            start_type="num", start_value=0,  start_color="F8D7DA",
            mid_type="num",   mid_value=50,   mid_color="FFF3CD",
            end_type="num",   end_value=100,  end_color="D4EDDA",
        ),
    )

# ── API Client ────────────────────────────────────────────────────────────────

class LWClient:
    def __init__(self):
        if not CLIENT_ID or not CLIENT_SECRET:
            raise ValueError(
                "Faltam credenciais. Configure LEARNWORLDS_CLIENT_ID e "
                "LEARNWORLDS_CLIENT_SECRET no arquivo .env"
            )
        self.base = BASE_URL.rstrip("/")
        self._token = None
        self._exp = 0

    def _auth(self):
        r = requests.post(
            f"{self.base}/v2/oauth2/access_token",
            headers={"Content-Type": "application/x-www-form-urlencoded",
                     "Lw-Client": CLIENT_ID},
            data={"client_id": CLIENT_ID, "client_secret": CLIENT_SECRET,
                  "grant_type": "client_credentials"},
        )
        r.raise_for_status()
        td = r.json().get("tokenData", r.json())
        self._token = td.get("token") or td.get("access_token")
        self._exp = time.time() + td.get("expires", td.get("expires_in", 3600)) - 60
        print("  ✓ Autenticado na LearnWorlds")

    def _h(self):
        if not self._token or time.time() >= self._exp:
            self._auth()
        return {"Authorization": f"Bearer {self._token}", "Lw-Client": CLIENT_ID}

    def get(self, path, params=None):
        r = requests.get(f"{self.base}{path}", headers=self._h(), params=params)
        r.raise_for_status()
        return r.json()

    def pages(self, path, page_size=25):
        params = {"page": 1, "itemsPerPage": page_size}
        while True:
            data = self.get(path, params)
            yield from data.get("data", [])
            meta = data.get("meta", {})
            if meta.get("page", 1) >= meta.get("totalPages", 1):
                break
            params["page"] += 1

# ── Extração de dados ─────────────────────────────────────────────────────────

def _days_ago(s):
    if not s:
        return None
    try:
        dt = datetime.fromisoformat(s.replace("Z", "+00:00"))
        return (datetime.now(timezone.utc) - dt).days
    except Exception:
        return None

def _fmt_date(s):
    if not s:
        return ""
    try:
        return datetime.fromisoformat(s.replace("Z", "+00:00")).strftime("%d/%m/%Y")
    except Exception:
        return s[:10] if len(s) >= 10 else s

def _parse_sections(prog):
    sections_out, quiz_scores = [], []
    quiz_total = quiz_done = video_total = video_done = 0
    for sec in prog.get("sections", []):
        units = sec.get("units", [])
        done_units = sum(1 for u in units if u.get("completed", False))
        sec_qs = []
        for u in units:
            t = u.get("type", "").lower()
            done = u.get("completed", False)
            score = u.get("score")
            if t in ("exam", "quiz", "assessment"):
                quiz_total += 1
                if done: quiz_done += 1
                if score is not None:
                    quiz_scores.append(score); sec_qs.append(score)
            elif t in ("video", "vimeo", "youtube"):
                video_total += 1
                if done: video_done += 1
        sections_out.append({
            "section_title": sec.get("title", "Módulo"),
            "total_units": len(units),
            "completed_units": done_units,
            "pct": round(done_units / len(units) * 100, 1) if units else 0,
        })
    stats = {
        "quiz_total": quiz_total, "quiz_done": quiz_done,
        "quiz_avg": round(sum(quiz_scores)/len(quiz_scores), 1) if quiz_scores else None,
        "video_done": video_done,
    }
    return sections_out, stats

def extract(client, course_id):
    print(f"\n📡 Buscando dados do curso: {course_id}")
    cr = client.get(f"/v2/courses/{course_id}")
    course = cr.get("data", cr)
    print(f"  ✓ Curso: {course.get('title', course_id)}")

    users = list(client.pages(f"/v2/courses/{course_id}/users"))
    print(f"  ✓ {len(users)} alunos matriculados")

    try:
        reviews = list(client.pages(f"/v2/courses/{course_id}/reviews"))
        print(f"  ✓ {len(reviews)} avaliações")
    except Exception:
        reviews = []

    records, section_records = [], []
    for i, user in enumerate(users, 1):
        uid = user.get("id") or user.get("userId")
        email = user.get("email", "")
        print(f"  [{i}/{len(users)}] {email}")
        try:
            resp = client.get(f"/v2/users/{uid}/course-progress/{course_id}")
            prog = resp.get("data", resp)
        except Exception:
            prog = {}

        enrolled_at   = user.get("enrolledAt") or user.get("created_at", "")
        last_activity = (prog.get("lastActivityAt") or prog.get("updatedAt")
                         or user.get("lastActivityAt", ""))
        comp_pct   = float(prog.get("completionPercentage", 0) or 0)
        comp_units = int(prog.get("completedUnits", 0) or 0)
        total_units = int(prog.get("totalUnits", 0) or 0)
        time_s     = int(prog.get("totalTimeSpentInSeconds", 0) or 0)
        status     = prog.get("status", "not_started") or "not_started"

        secs, ustats = _parse_sections(prog)
        quiz_avg = ustats["quiz_avg"] or prog.get("score")

        records.append({
            "email": email,
            "username": user.get("username") or user.get("name") or "",
            "enrolled_at_fmt": _fmt_date(enrolled_at),
            "last_activity_fmt": _fmt_date(last_activity),
            "days_since_enrollment": _days_ago(enrolled_at),
            "days_since_last_access": _days_ago(last_activity),
            "completion_pct": comp_pct,
            "completed_lessons": comp_units,
            "total_lessons": total_units,
            "pct_lessons": round(comp_units/total_units*100, 1) if total_units else 0,
            "time_h": round(time_s/3600, 2),
            "time_min": round(time_s/60, 1),
            "status": status,
            "status_pt": STATUS_PT.get(status, "Desconhecido"),
            "quiz_score": quiz_avg,
            "quiz_done": ustats["quiz_done"],
            "quiz_total": ustats["quiz_total"],
            "video_done": ustats["video_done"],
            "certificate": bool(prog.get("certificateIssued", False)),
        })
        for sec in secs:
            section_records.append({"email": email, **sec})

    return course, records, section_records, reviews

# ── Construtores de abas ──────────────────────────────────────────────────────

def build_resumo(ws, course, df, reviews):
    _col_widths(ws, {"A": 32, "B": 18, "C": 14})
    total     = len(df)
    completed = int((df["completion_pct"] >= 100).sum())
    in_prog   = int(((df["completion_pct"] > 0) & (df["completion_pct"] < 100)).sum())
    not_start = int((df["completion_pct"] == 0).sum())
    comp_rate = round(completed/total*100, 1) if total else 0
    avg_comp  = round(df["completion_pct"].mean(), 1) if total else 0
    avg_h     = round(df["time_h"].mean(), 1) if total else 0
    total_h   = round(df["time_h"].sum(), 1) if total else 0
    certs     = int(df["certificate"].sum())
    avg_quiz  = round(df["quiz_score"].dropna().mean(), 1) if df["quiz_score"].notna().any() else None
    avg_rat   = None
    if reviews:
        scores = [r.get("rating") for r in reviews if r.get("rating") is not None]
        if scores: avg_rat = round(sum(scores)/len(scores), 1)
    now = datetime.now().strftime("%d/%m/%Y às %H:%M")

    # Título
    ws.merge_cells("A1:C1")
    c = ws.cell(row=1, column=1, value=course.get("title", "Curso"))
    c.fill = fill(NAVY); c.font = font(WHITE, bold=True, size=14)
    c.alignment = align(h="center"); ws.row_dimensions[1].height = 40
    ws.cell(row=2, column=1, value=f"Relatório gerado em: {now}").font = font("888888", italic=True, size=9)

    # Blocos de KPIs
    kpis = [
        (4,  "MÉTRICAS GERAIS", [
            ("Total de Alunos Matriculados", total,            None),
            ("Taxa de Conclusão",            f"{comp_rate}%",  None),
            ("Concluíram",                   completed,        GREEN_BG),
            ("Em Andamento",                 in_prog,          AMBER_BG),
            ("Não Iniciaram",               not_start,         RED_BG),
        ]),
        (11, "PERFORMANCE", [
            ("Progresso Médio",              f"{avg_comp}%",   None),
            ("Tempo Médio por Aluno",        f"{avg_h}h",      None),
            ("Tempo Total da Turma",         f"{total_h}h",    None),
            ("Certificados Emitidos",        certs,            None),
            ("Média dos Quizzes",           f"{avg_quiz}%" if avg_quiz else "—", None),
            ("Satisfação Média",            f"{avg_rat}/5" if avg_rat else "—",  None),
        ]),
    ]
    for start, title, items in kpis:
        ws.merge_cells(f"A{start}:C{start}")
        c = ws.cell(row=start, column=1, value=title)
        c.fill = fill(NAVY); c.font = font(WHITE, bold=True)
        c.alignment = align(h="left"); ws.row_dimensions[start].height = 20
        for j, (label, val, row_bg) in enumerate(items, 1):
            r = start + j
            ws.row_dimensions[r].height = 18
            ca = ws.cell(row=r, column=1, value=label)
            cb = ws.cell(row=r, column=2, value=val)
            for cell in (ca, cb):
                cell.font = font(size=10)
                cell.fill = fill(row_bg if row_bg else (LT_GRAY if j % 2 == 0 else WHITE))
            ca.alignment = align(h="left"); cb.alignment = align(h="center")

    # Tabela de status
    _header_row(ws, 19, ["DISTRIBUIÇÃO DE STATUS", "Qtd", "% do Total"], bg="2D6A9F")
    for j, (label, qty, pct, bg_c) in enumerate([
        ("Concluído",    completed,  comp_rate,                                             GREEN_BG),
        ("Em Andamento", in_prog,   round(in_prog/total*100,   1) if total else 0,         AMBER_BG),
        ("Não Iniciado", not_start, round(not_start/total*100, 1) if total else 0,         RED_BG),
    ], 20):
        for col, val in enumerate([label, qty, f"{pct}%"], 1):
            c = ws.cell(row=j, column=col, value=val)
            c.fill = fill(bg_c); c.font = font(size=10)
            c.alignment = align(h="center")


def build_alunos(ws, df):
    headers = ["Nome", "Email", "Status", "Progresso (%)", "Aulas Concluídas",
               "Total Aulas", "% Aulas", "Tempo (h)", "Quizzes Feitos",
               "Score Quiz (%)", "Certificado", "Dias Matriculado",
               "Último Acesso", "Dias s/ Acesso"]
    _header_row(ws, 1, headers)
    ws.freeze_panes = "A2"
    _col_widths(ws, {"A":20,"B":30,"C":16,"D":13,"E":14,"F":12,"G":12,
                     "H":12,"I":14,"J":14,"K":13,"L":15,"M":15,"N":15})
    df_s = df.sort_values("completion_pct", ascending=False)
    for i, r in enumerate(df_s.to_dict("records"), 2):
        bg = LT_GRAY if i % 2 == 0 else WHITE
        st_bg, st_fg = _status_fill(r["status_pt"])
        vals = [r["username"], r["email"], r["status_pt"],
                r["completion_pct"], r["completed_lessons"], r["total_lessons"],
                r["pct_lessons"], r["time_h"], r.get("quiz_done", 0),
                r.get("quiz_score") or "", "Sim" if r["certificate"] else "Não",
                r.get("days_since_enrollment") or "",
                r["last_activity_fmt"], r.get("days_since_last_access") or ""]
        for col, val in enumerate(vals, 1):
            c = ws.cell(row=i, column=col, value=val)
            c.font = font(size=10); c.fill = fill(bg)
            c.alignment = align(h="center" if col > 2 else "left")
        ws.cell(row=i, column=3).fill = fill(st_bg)
        ws.cell(row=i, column=3).font = font(st_fg, bold=True)
    n = len(df_s)
    if n > 0:
        _color_scale(ws, 2, n+1, "D")
        _color_scale(ws, 2, n+1, "G")


def build_presenca(ws, df):
    headers = ["Nome", "Email", "Status", "Data Matrícula", "Último Acesso",
               "Dias s/ Acesso", "Dias Matriculado", "Tempo (h)", "Progresso (%)", "Situação de Acesso"]
    _header_row(ws, 1, headers)
    ws.freeze_panes = "A2"
    _col_widths(ws, {"A":20,"B":30,"C":16,"D":14,"E":14,"F":14,"G":15,"H":12,"I":13,"J":20})
    df_s = df.sort_values("days_since_last_access", ascending=False, na_position="first")
    for i, r in enumerate(df_s.to_dict("records"), 2):
        label, ac_bg, ac_fg = _access_label(r.get("days_since_last_access"))
        st_bg, st_fg = _status_fill(r["status_pt"])
        bg = LT_GRAY if i % 2 == 0 else WHITE
        vals = [r["username"], r["email"], r["status_pt"],
                r["enrolled_at_fmt"], r["last_activity_fmt"],
                r.get("days_since_last_access") or "—",
                r.get("days_since_enrollment") or "—",
                r["time_h"], r["completion_pct"], label]
        for col, val in enumerate(vals, 1):
            c = ws.cell(row=i, column=col, value=val)
            c.font = font(size=10); c.fill = fill(bg)
            c.alignment = align(h="center" if col > 2 else "left")
        ws.cell(row=i, column=3).fill = fill(st_bg)
        ws.cell(row=i, column=3).font = font(st_fg, bold=True)
        ws.cell(row=i, column=10).fill = fill(ac_bg)
        ws.cell(row=i, column=10).font = font(ac_fg, bold=True)
    n = len(df_s)
    if n > 0:
        _color_scale(ws, 2, n+1, "I")


def build_aulas_quiz(ws, df):
    headers = ["Nome", "Email", "Aulas Concluídas", "Total Aulas", "% Aulas",
               "Vídeos Assistidos", "Quizzes Feitos", "Total Quizzes",
               "Score Quiz (%)", "Tempo (h)", "Tempo/Aula (min)", "Status"]
    _header_row(ws, 1, headers)
    ws.freeze_panes = "A2"
    _col_widths(ws, {"A":20,"B":30,"C":14,"D":12,"E":12,"F":15,"G":14,"H":14,"I":14,"J":12,"K":16,"L":16})
    df_s = df.sort_values("completed_lessons", ascending=False)
    for i, r in enumerate(df_s.to_dict("records"), 2):
        bg = LT_GRAY if i % 2 == 0 else WHITE
        min_per = round(r["time_min"] / r["completed_lessons"], 1) if r["completed_lessons"] > 0 else 0
        vals = [r["username"], r["email"],
                r["completed_lessons"], r["total_lessons"], r["pct_lessons"],
                r.get("video_done", 0), r.get("quiz_done", 0), r.get("quiz_total", 0),
                r.get("quiz_score") or "", r["time_h"], min_per, r["status_pt"]]
        for col, val in enumerate(vals, 1):
            c = ws.cell(row=i, column=col, value=val)
            c.font = font(size=10); c.fill = fill(bg)
            c.alignment = align(h="center" if col > 2 else "left")
        st_bg, st_fg = _status_fill(r["status_pt"])
        ws.cell(row=i, column=12).fill = fill(st_bg)
        ws.cell(row=i, column=12).font = font(st_fg, bold=True)
    n = len(df_s)
    if n > 0:
        _color_scale(ws, 2, n+1, "E")
        _color_scale(ws, 2, n+1, "I")


def build_modulos(ws, section_records):
    if not section_records:
        ws.cell(row=1, column=1, value="Dados de módulos não disponíveis via API.").font = font(italic=True)
        ws.column_dimensions["A"].width = 50
        return
    df_s = pd.DataFrame(section_records)
    pivot = df_s.pivot_table(index="email", columns="section_title",
                              values="pct", fill_value=0).reset_index()
    sec_cols = [c for c in pivot.columns if c != "email"]
    pivot["Média (%)"] = pivot[sec_cols].mean(axis=1).round(1)
    all_cols = ["email"] + sec_cols + ["Média (%)"]
    _header_row(ws, 1, all_cols)
    ws.freeze_panes = "B2"
    ws.column_dimensions["A"].width = 30
    for i, col in enumerate(sec_cols + ["Média (%)"], 2):
        ws.column_dimensions[get_column_letter(i)].width = max(16, len(str(col)) * 1.1)
    for i, r in enumerate(pivot.to_dict("records"), 2):
        bg = LT_GRAY if i % 2 == 0 else WHITE
        for col, key in enumerate(all_cols, 1):
            c = ws.cell(row=i, column=col, value=r.get(key, ""))
            c.font = font(size=10); c.fill = fill(bg)
            c.alignment = align(h="center" if col > 1 else "left")
    n = len(pivot)
    if n > 0:
        for col_i in range(2, len(all_cols) + 1):
            _color_scale(ws, 2, n+1, get_column_letter(col_i))


def build_em_risco(ws, df):
    at_risk = df[
        (df["completion_pct"] < 25) |
        (df["days_since_last_access"].fillna(999) > 14)
    ].sort_values("completion_pct").copy()

    if at_risk.empty:
        ws.cell(row=1, column=1,
                value="✅ Nenhum aluno em risco! Todos estão engajados.").font = font(GREEN_FG, bold=True, size=11)
        ws.column_dimensions["A"].width = 50
        return

    headers = ["Nome", "Email", "Status", "Progresso (%)", "Dias s/ Acesso",
               "Último Acesso", "Tempo (h)", "Motivo de Atenção"]
    _header_row(ws, 1, headers, bg="B71C1C")
    ws.freeze_panes = "A2"
    _col_widths(ws, {"A":20,"B":30,"C":16,"D":13,"E":14,"F":14,"G":12,"H":38})

    def _motivo(r):
        parts = []
        if r["completion_pct"] < 25:
            parts.append(f"Progresso baixo ({r['completion_pct']:.0f}%)")
        dsa = r.get("days_since_last_access")
        if dsa is not None and not (isinstance(dsa, float) and dsa != dsa) and dsa > 14:
            parts.append(f"Sem acesso há {int(dsa)} dias")
        if (dsa is None or (isinstance(dsa, float) and dsa != dsa)) and r["completion_pct"] == 0:
            parts.append("Nunca acessou")
        return " | ".join(parts) if parts else "—"

    for i, r in enumerate(at_risk.to_dict("records"), 2):
        st_bg, st_fg = _status_fill(r["status_pt"])
        vals = [r["username"], r["email"], r["status_pt"], r["completion_pct"],
                int(r["days_since_last_access"]) if pd.notna(r.get("days_since_last_access")) else "—",
                r["last_activity_fmt"], r["time_h"], _motivo(r)]
        for col, val in enumerate(vals, 1):
            c = ws.cell(row=i, column=col, value=val)
            c.font = font(size=10); c.fill = fill(RED_BG)
            c.alignment = align(h="center" if col not in (1, 2, 8) else "left", wrap=(col == 8))
        ws.cell(row=i, column=3).fill = fill(st_bg)
        ws.cell(row=i, column=3).font = font(st_fg, bold=True)


def build_satisfacao(ws, reviews):
    if not reviews:
        ws.cell(row=1, column=1, value="Nenhuma avaliação registrada para este curso.").font = font(italic=True)
        ws.column_dimensions["A"].width = 50
        return
    df_r = pd.DataFrame(reviews)
    if "rating" not in df_r.columns or df_r["rating"].dropna().empty:
        ws.cell(row=1, column=1, value="Dados de satisfação não disponíveis via API.").font = font(italic=True)
        return
    avg = round(df_r["rating"].dropna().mean(), 2)
    dist = df_r["rating"].value_counts().sort_index()
    _header_row(ws, 1, ["Avaliação Média", str(avg)], bg="1565C0")
    _header_row(ws, 2, ["Total de Avaliações", str(len(df_r))], bg="1565C0")
    _header_row(ws, 4, ["Nota", "Qtd", "% do Total"])
    for j, (k, v) in enumerate(dist.items(), 5):
        bg = GREEN_BG if k >= 4 else AMBER_BG if k >= 3 else RED_BG
        for col, val in enumerate([int(k), int(v), f"{round(v/len(df_r)*100,1)}%"], 1):
            c = ws.cell(row=j, column=col, value=val)
            c.fill = fill(bg); c.font = font(size=10); c.alignment = align(h="center")
    start = 5 + len(dist) + 2
    _header_row(ws, start, ["Email/ID", "Nota", "Comentário", "Data"])
    for j, r in enumerate(reviews, start+1):
        for col, val in enumerate([r.get("user_id",""), r.get("rating",""),
                                   r.get("comment",""), r.get("created_at","")], 1):
            ws.cell(row=j, column=col, value=val).font = font(size=10)
    _col_widths(ws, {"A":30,"B":10,"C":50,"D":15})

# ── Main ───────────────────────────────────────────────────────────────────────

def run(course_id: str):
    client = LWClient()
    course, records, section_records, reviews = extract(client, course_id)

    if not records:
        raise ValueError("Nenhum aluno encontrado no curso.")

    df = pd.DataFrame(records)
    for col in ["completion_pct","time_h","time_min","completed_lessons",
                "total_lessons","days_since_last_access","days_since_enrollment"]:
        df[col] = pd.to_numeric(df[col], errors="coerce").fillna(0)

    wb = Workbook()
    sheets = [
        ("📊 Resumo Geral",     lambda ws: build_resumo(ws, course, df, reviews)),
        ("👥 Roster de Alunos",  lambda ws: build_alunos(ws, df)),
        ("📅 Presença & Acesso", lambda ws: build_presenca(ws, df)),
        ("📝 Aulas & Quizzes",   lambda ws: build_aulas_quiz(ws, df)),
        ("🔥 Por Módulo",        lambda ws: build_modulos(ws, section_records)),
        ("⚠️ Em Risco",          lambda ws: build_em_risco(ws, df)),
        ("⭐ Satisfação",        lambda ws: build_satisfacao(ws, reviews)),
    ]

    wb.active.title = sheets[0][0]
    sheets[0][1](wb.active)
    for name, builder in sheets[1:]:
        builder(wb.create_sheet(name))

    out = f"relatorio_{course_id}_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    wb.save(out)
    return out


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    print("\n🚀 LearnWorlds → Excel Report")
    print("=" * 42)
    path = run(sys.argv[1])
    print(f"\n✅ Arquivo gerado: {path}")
    print("   Suba no Google Sheets: sheets.new → Arquivo → Importar\n")
